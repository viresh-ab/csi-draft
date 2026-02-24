"""
Tests for document_loader.py

Run with: pytest tests/test_document_loader.py -v
"""
import pytest
from pathlib import Path
from unittest.mock import patch, MagicMock
from app.retrieval.document_loader import resolve_path, extract_text, load_case_study


@patch("app.retrieval.document_loader.CS_ROOT", Path("/fake/cs-files"))
def test_resolve_path_not_found_raises():
    with pytest.raises(FileNotFoundError):
        resolve_path("BFSI", "nonexistent_file.pdf")


def test_extract_text_unsupported_format(tmp_path):
    fake_file = tmp_path / "test.xlsx"
    fake_file.write_text("fake content")
    with pytest.raises(ValueError, match="Unsupported file type"):
        extract_text(fake_file)


def test_extract_text_docx(tmp_path):
    from docx import Document
    doc = Document()
    doc.add_paragraph("This is a test paragraph.")
    doc.add_paragraph("Second paragraph.")
    docx_path = tmp_path / "test.docx"
    doc.save(str(docx_path))

    text = extract_text(docx_path)
    assert "test paragraph" in text
    assert "Second paragraph" in text


def test_extract_text_pptx(tmp_path):
    from pptx import Presentation
    prs = Presentation()
    slide_layout = prs.slide_layouts[5]
    slide = prs.slides.add_slide(slide_layout)
    txBox = slide.shapes.add_textbox(0, 0, 100, 50)
    txBox.text_frame.text = "Hello from PPTX"
    pptx_path = tmp_path / "test.pptx"
    prs.save(str(pptx_path))

    text = extract_text(pptx_path)
    assert "Hello from PPTX" in text
